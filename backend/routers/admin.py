from fastapi import APIRouter, Depends, HTTPException, Query, status, UploadFile, File
from typing import Optional, List, Dict, Any
import csv
import io
from pydantic import BaseModel, Field
from datetime import datetime, timezone
from supabase import create_client

from core.config import settings
from core.security import AuthenticatedUser, require_role, get_current_user
from core import clock
from services import scheduler
from services import audit as audit_service
from services import approvals

router = APIRouter(prefix="/admin", tags=["Admin"])
_supabase = create_client(settings.SUPABASE_URL, settings.SUPABASE_SERVICE_KEY)


class ClockAdvanceRequest(BaseModel):
    days: float = Field(..., description="Number of days to advance (can be fractional or negative)")


class SLARuleUpdate(BaseModel):
    id: Optional[str] = None
    region_id: Optional[str] = None
    scope: str
    stage_key: Optional[str] = None
    template_code: Optional[str] = None
    duration_days: int
    reminder_offsets_days: List[int] = [7, 3, 1]
    is_active: bool = True


class TemplateUpdate(BaseModel):
    id: Optional[str] = None
    code: str
    name: str
    definition: Dict[str, Any]
    is_active: bool = True


class ApprovalDecision(BaseModel):
    reason: Optional[str] = None


class SettingsUpdate(BaseModel):
    maker_checker_for_overrides: bool


class UserUpdate(BaseModel):
    role: Optional[str] = Field(None, pattern="^(institution|scrutiny_officer|regional_officer|evc_member|sac_shc_member|bureau_officer|auditor|admin)$")
    region_id: Optional[str] = None
    institution_id: Optional[str] = None
    full_name: Optional[str] = None
    locale: Optional[str] = Field(None, pattern="^(en|hi)$")


CHANNELS = {"in_app", "email", "sms", "whatsapp"}


class EscalationChainUpdate(BaseModel):
    id: Optional[str] = None
    region_id: Optional[str] = None
    scope: str
    levels: List[Dict[str, Any]]
    ack_timeout_hours: int = 48


# -------------------------------------------------------------
# Time-Travel Simulator Endpoints (§5.1, §6, §7.1)
# -------------------------------------------------------------
@router.get("/clock")
async def get_clock(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    """Get current simulated clock status."""
    sim_now = clock.now()
    actual_utc = datetime.now(timezone.utc)
    offset_days = clock.get_offset_days()
    return {
        "clock_now": sim_now.isoformat(),
        "actual_utc": actual_utc.isoformat(),
        "offset_seconds": int(offset_days * 86400),
        "offset_days": round(offset_days, 2),
        "is_simulated": abs(offset_days) > 0.0001
    }


@router.post("/clock/advance")
async def advance_clock(
    req: ClockAdvanceRequest,
    user: AuthenticatedUser = Depends(require_role(["admin"]))
):
    """
    Advance simulated time by N days and run scheduler immediately (§5.1).
    Produces reminders, overdue alerts, and escalations accordingly.
    """
    before_time = clock.now().isoformat()
    new_time = clock.advance_days(req.days)

    # Run scheduler immediately after clock advance
    run_result = scheduler.run_once()

    # Record audit log
    audit_service.record(
        actor=user,
        action="CLOCK_ADVANCE",
        entity_type="clock",
        entity_id=None,
        before={"simulated_time": before_time},
        after={"simulated_time": new_time.isoformat(), "days_advanced": req.days},
        reason=f"Admin advanced clock by {req.days} days"
    )

    return {
        "ok": True,
        "new_clock_now": new_time.isoformat(),
        "offset_days": round(clock.get_offset_days(), 2),
        "scheduler_run": run_result
    }


@router.post("/clock/reset")
async def reset_clock(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    """Reset simulated clock back to real-time UTC and trigger scheduler (§5.1)."""
    before_time = clock.now().isoformat()
    clock.reset()
    reset_time = clock.now()

    run_result = scheduler.run_once()

    audit_service.record(
        actor=user,
        action="CLOCK_RESET",
        entity_type="clock",
        entity_id=None,
        before={"simulated_time": before_time},
        after={"simulated_time": reset_time.isoformat(), "offset_days": 0},
        reason="Admin reset clock to real UTC"
    )

    return {
        "ok": True,
        "new_clock_now": reset_time.isoformat(),
        "offset_days": 0,
        "scheduler_run": run_result
    }


# -------------------------------------------------------------
# Scheduler Management Endpoints (§5.4)
# -------------------------------------------------------------
@router.get("/scheduler/runs")
async def list_scheduler_runs(
    limit: int = Query(20, ge=1, le=100),
    user: AuthenticatedUser = Depends(require_role(["admin"]))
):
    """List execution history of the scheduler runs."""
    res = _supabase.table("scheduler_runs").select("*").order("started_at", desc=True).limit(limit).execute()
    return res.data or []


@router.post("/scheduler/run")
async def trigger_scheduler_run(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    """Manually trigger a full scheduler pass immediately."""
    result = scheduler.run_once()
    return {
        "ok": True,
        "run": result
    }


# -------------------------------------------------------------
# SLA Rules & Escalation Chains (§6)
# -------------------------------------------------------------
@router.get("/sla-rules")
async def get_sla_rules(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    """List all configured SLA rules."""
    res = _supabase.table("sla_rules").select("*").order("scope").execute()
    return res.data or []


def _request_rule_change(table: str, row: Dict[str, Any], user: AuthenticatedUser) -> Dict[str, Any]:
    """Rule changes never apply directly: a second administrator approves them (§15.2)."""
    request = approvals.create_request(
        "rule_change", table, row.get("id"), {"table": table, "row": row}, user,
        reason=f"Change to {table}",
    )
    return {"status": "pending_approval", "approval_request": request}


@router.put("/sla-rules", status_code=202)
async def update_sla_rule(
    rule: SLARuleUpdate,
    user: AuthenticatedUser = Depends(require_role(["admin"]))
):
    if rule.duration_days <= 0 or any(o <= 0 for o in rule.reminder_offsets_days):
        raise HTTPException(status_code=422, detail="Durations and reminder offsets must be positive")
    return _request_rule_change("sla_rules", rule.model_dump(), user)


@router.get("/escalation-chains")
async def get_escalation_chains(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    res = _supabase.table("escalation_chains").select("*").execute()
    return res.data or []


@router.put("/escalation-chains", status_code=202)
async def update_escalation_chain(
    chain: EscalationChainUpdate,
    user: AuthenticatedUser = Depends(require_role(["admin"]))
):
    levels = sorted(chain.levels, key=lambda l: l.get("level", 0))
    for lvl in levels:
        if not isinstance(lvl.get("level"), int) or not isinstance(lvl.get("after_days_overdue"), (int, float)) or not lvl.get("notify_role"):
            raise HTTPException(status_code=422, detail="Each level needs level, after_days_overdue and notify_role")
        bad = set(lvl.get("channels") or ["in_app"]) - CHANNELS
        if bad:
            raise HTTPException(status_code=422, detail=f"Unknown channel(s): {', '.join(sorted(bad))}")
    payload = chain.model_dump()
    payload["levels"] = levels
    return _request_rule_change("escalation_chains", payload, user)


# -------------------------------------------------------------
# Workflow templates (§4, §6)
# -------------------------------------------------------------
@router.get("/templates")
async def get_templates(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    return _supabase.table("workflow_templates").select("*").order("code").execute().data or []


@router.put("/templates", status_code=202)
async def update_template(
    tmpl: TemplateUpdate,
    user: AuthenticatedUser = Depends(require_role(["admin"]))
):
    stages = tmpl.definition.get("stages") or {}
    start = tmpl.definition.get("start")
    if start not in stages:
        raise HTTPException(status_code=422, detail="definition.start must name a stage")
    for key, st in stages.items():
        for ref in ("next", "on_return"):
            if st.get(ref) and st[ref] not in stages:
                raise HTTPException(status_code=422, detail=f"Stage '{key}' {ref} points to unknown stage '{st[ref]}'")
        if not st.get("owner_role") or not isinstance(st.get("sla_days"), int) or st["sla_days"] <= 0:
            raise HTTPException(status_code=422, detail=f"Stage '{key}' needs an owner_role and positive sla_days")
    return _request_rule_change("workflow_templates", tmpl.model_dump(exclude_none=True), user)


# -------------------------------------------------------------
# Maker-checker queue (§15.2)
# -------------------------------------------------------------
@router.get("/approvals")
async def list_approvals(
    status_filter: Optional[str] = Query(None, alias="status"),
    user: AuthenticatedUser = Depends(require_role(["admin", "bureau_officer"]))
):
    """Administrators see rule changes; bureau officers see reverts and overrides."""
    kinds = ["rule_change"] if user.role == "admin" else ["revert", "assurance_override"]
    q = _supabase.table("approval_requests").select("*").in_("kind", kinds)
    if status_filter:
        q = q.eq("status", status_filter)
    rows = q.order("created_at", desc=True).limit(100).execute().data or []
    names = {p["user_id"]: p["full_name"] for p in (_supabase.table("profiles").select("user_id, full_name").execute().data or [])}
    for r in rows:
        r["requested_by_name"] = names.get(r["requested_by"])
        r["approved_by_name"] = names.get(r.get("approved_by"))
        r["can_decide"] = r["status"] == "pending" and str(r["requested_by"]) != str(user.id)
    return rows


@router.post("/approvals/{request_id}/{decision}")
async def decide_approval(
    request_id: str,
    decision: str,
    body: ApprovalDecision = ApprovalDecision(),
    user: AuthenticatedUser = Depends(require_role(["admin", "bureau_officer"]))
):
    if decision not in ("approve", "reject"):
        raise HTTPException(status_code=404, detail="Unknown decision")
    try:
        return await approvals.decide(request_id, user, decision == "approve", body.reason)
    except PermissionError as e:
        raise HTTPException(status_code=403, detail=str(e))
    except (approvals.ApprovalError, ValueError) as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/settings")
async def get_settings(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    return {"maker_checker_for_overrides": approvals.overrides_need_approval()}


@router.put("/settings")
async def update_settings(body: SettingsUpdate, user: AuthenticatedUser = Depends(require_role(["admin"]))):
    before = approvals.overrides_need_approval()
    approvals.set_setting("maker_checker_for_overrides", body.maker_checker_for_overrides)
    audit_service.record(
        actor=user, action="SETTING_CHANGED", entity_type="system_settings", entity_id=None,
        before={"maker_checker_for_overrides": before},
        after={"maker_checker_for_overrides": body.maker_checker_for_overrides},
        reason="Maker-checker for overrides and reverts",
    )
    return {"maker_checker_for_overrides": body.maker_checker_for_overrides}


# -------------------------------------------------------------
# User management (§15.0); contact details masked in lists (§15.8)
# -------------------------------------------------------------
def mask_email(email: Optional[str]) -> Optional[str]:
    if not email or "@" not in email:
        return email
    local, domain = email.split("@", 1)
    return f"{local[0]}***@{domain}"


def mask_phone(phone: Optional[str]) -> Optional[str]:
    if not phone:
        return phone
    digits = [c for c in phone if c.isdigit()]
    return "******" + "".join(digits[-4:]) if len(digits) >= 4 else "****"


def _auth_emails() -> Dict[str, str]:
    emails: Dict[str, str] = {}
    page = 1
    while True:
        users = _supabase.auth.admin.list_users(page=page, per_page=200)
        for u in users:
            emails[u.id] = u.email
        if len(users) < 200:
            return emails
        page += 1


@router.get("/users")
async def list_users(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    emails = _auth_emails()
    profiles = _supabase.table("profiles").select("*, regions(name), institutions(name)").order("role").execute().data or []
    return [{
        "user_id": p["user_id"],
        "full_name": p["full_name"],
        "role": p["role"],
        "region": (p.get("regions") or {}).get("name"),
        "institution": (p.get("institutions") or {}).get("name"),
        "email": mask_email(emails.get(p["user_id"])),
        "phone": mask_phone(p.get("phone")),
        "locale": p.get("locale"),
    } for p in profiles]


@router.get("/users/{user_id}")
async def get_user(user_id: str, user: AuthenticatedUser = Depends(require_role(["admin"]))):
    """Detail view reveals full contact details; the reveal is audited."""
    prof = _supabase.table("profiles").select("*").eq("user_id", user_id).limit(1).execute().data
    if not prof:
        raise HTTPException(status_code=404, detail="User not found")
    auth_user = _supabase.auth.admin.get_user_by_id(user_id).user
    audit_service.record(
        actor=user, action="USER_CONTACT_VIEWED", entity_type="profile", entity_id=user_id,
        before=None, after=None, reason="Administrator opened user detail",
    )
    return {**prof[0], "email": auth_user.email if auth_user else None}


@router.patch("/users/{user_id}")
async def update_user(user_id: str, body: UserUpdate, user: AuthenticatedUser = Depends(require_role(["admin"]))):
    prof = _supabase.table("profiles").select("*").eq("user_id", user_id).limit(1).execute().data
    if not prof:
        raise HTTPException(status_code=404, detail="User not found")
    if user_id == user.id and body.role and body.role != "admin":
        raise HTTPException(status_code=400, detail="Administrators cannot change their own role")
    changes = body.model_dump(exclude_none=True)
    updated = _supabase.table("profiles").update(changes).eq("user_id", user_id).execute().data[0]
    # Keep the token claims in step with the profile
    _supabase.auth.admin.update_user_by_id(user_id, {"app_metadata": {
        "role": updated["role"], "region_id": updated.get("region_id"), "institution_id": updated.get("institution_id"),
    }})
    audit_service.record(
        actor=user, action="USER_UPDATED", entity_type="profile", entity_id=user_id,
        before={k: prof[0].get(k) for k in changes}, after=changes, reason="Administrator updated user",
    )
    return updated


# -------------------------------------------------------------
# Bulk import (§5.6): CSV or Excel for institutions, applications and documents
# -------------------------------------------------------------
IMPORT_COLUMNS = {
    "institutions": {"required": ["name", "aicte_id"], "optional": ["state", "city", "region", "kind", "contact_email"]},
    "applications": {"required": ["application_no", "aicte_id", "application_type", "template_code"], "optional": []},
    "documents": {"required": ["aicte_id", "doc_type", "title"], "optional": ["valid_from", "valid_until", "application_no"]},
}


def _read_rows(filename: str, content: bytes) -> List[Dict[str, Any]]:
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in reader]
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(content), read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip() for h in rows[0]]
        return [{header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r)} for r in rows[1:] if any(r)]
    raise HTTPException(status_code=422, detail="Upload a .csv or .xlsx file")


@router.post("/import")
async def bulk_import(
    entity: str = Query(..., pattern="^(institutions|applications|documents)$"),
    file: UploadFile = File(...),
    user: AuthenticatedUser = Depends(require_role(["admin"]))
):
    """
    Import rows through the mock adapter path. Applications start their workflow through the
    workflow engine as the institution would; every row is validated and errors are reported per row.
    """
    content = await file.read()
    if len(content) > settings.MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File too large")
    rows = _read_rows(file.filename or "", content)
    spec = IMPORT_COLUMNS[entity]
    regions = {r["name"].lower(): r["id"] for r in (_supabase.table("regions").select("id, name").execute().data or [])}
    institutions = {i["aicte_id"]: i for i in (_supabase.table("institutions").select("id, aicte_id, region_id").execute().data or [])}

    created, errors = 0, []
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        missing = [c for c in spec["required"] if not row.get(c)]
        if missing:
            errors.append({"row": n, "error": f"missing {', '.join(missing)}"})
            continue
        try:
            if entity == "institutions":
                region_id = regions.get((row.get("region") or "").lower())
                if row.get("region") and not region_id:
                    raise ValueError(f"unknown region '{row['region']}'")
                if row["aicte_id"] in institutions:
                    raise ValueError("aicte_id already exists")
                _supabase.table("institutions").insert({
                    "name": row["name"], "aicte_id": row["aicte_id"], "state": row.get("state") or None,
                    "city": row.get("city") or None, "region_id": region_id,
                    "kind": row.get("kind") or "existing", "contact_email": row.get("contact_email") or None,
                }).execute()
            elif entity == "applications":
                inst = institutions.get(row["aicte_id"])
                if not inst:
                    raise ValueError(f"unknown institution '{row['aicte_id']}'")
                from services.workflow_engine import engine
                app = engine.start(inst["id"], row["application_type"], row["template_code"], user, region_id=inst.get("region_id"))
                _supabase.table("applications").update({"application_no": row["application_no"]}).eq("id", app["id"]).execute()
            else:
                inst = institutions.get(row["aicte_id"])
                if not inst:
                    raise ValueError(f"unknown institution '{row['aicte_id']}'")
                app_id = None
                if row.get("application_no"):
                    a = _supabase.table("applications").select("id").eq("application_no", row["application_no"]).limit(1).execute().data
                    if not a:
                        raise ValueError(f"unknown application '{row['application_no']}'")
                    app_id = a[0]["id"]
                for d in ("valid_from", "valid_until"):
                    if row.get(d):
                        datetime.strptime(row[d][:10], "%Y-%m-%d")
                _supabase.table("documents").insert({
                    "institution_id": inst["id"], "application_id": app_id, "doc_type": row["doc_type"],
                    "title": row["title"], "valid_from": (row.get("valid_from") or "")[:10] or None,
                    "valid_until": (row.get("valid_until") or "")[:10] or None, "status": "active",
                }).execute()
            created += 1
        except Exception as e:
            errors.append({"row": n, "error": str(e)})

    audit_service.record(
        actor=user, action="BULK_IMPORT", entity_type=entity, entity_id=None,
        before=None, after={"file": file.filename, "created": created, "errors": len(errors)},
        reason=f"Bulk import of {entity}",
    )
    return {"entity": entity, "rows": len(rows), "created": created, "errors": errors}


# -------------------------------------------------------------
# Admin Seed Endpoint (§6, §9)
# -------------------------------------------------------------
@router.post("/seed")
async def run_seed_cases(user: AuthenticatedUser = Depends(require_role(["admin"]))):
    """Seed or reseed demo cases (Section 9)."""
    from seed.demo_cases import seed_scripted_cases
    cases = seed_scripted_cases()
    return {
        "ok": True,
        "message": "Scripted demo cases seeded successfully",
        "cases": cases
    }
